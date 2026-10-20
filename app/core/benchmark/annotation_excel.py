"""K-150 etiketleme için Excel formu üretimi ve geri aktarımı.

Patron md tablosunda ÇALIŞAMAZ — Excel formu insan yüzüdür. Tek-kaynak
sözleşmesi DEĞİŞMEZ: kanonik insan girdisi yine `annotation_sheet*.md`,
JSON'a aktarım yine `scripts/import_annotations.py`. Bu modül yalnız iki
uç ekler:

  build  : md/json'daki güncel durumdan Excel formu üretir (kısmi etiket
           varsa taşınır — yeniden üretim idempotenttir). Form: açılır
           listeli hücreler, boş-hücre kırmızı vurgusu, kilitli kimlik/
           kelime sütunları, TALİMAT + TANIMLAR + ÖRNEKLER sekmeleri.
  import : doldurulmuş Excel'i doğrulayıp md tablosunu YENİDEN YAZAR
           (tablo üstündeki başlık metni korunur). Değerler şemaya göre
           fail-closed doğrulanır; kelime/kimlik oynanmışsa reddedilir.

İkinci etiketçi formu AYRI dosyadır ve birincil etiketleri İÇERMEZ
(körlük sözleşmesi, annotation_guide.md).
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from app.core.benchmark.annotation_io import (
    AnnotationError,
    parse_primary_md,
    parse_second_md,
)

LABEL_FIELDS = ["intent_type", "gt", "ga", "ads_class", "opinion_discussion",
                "curiosity_comparison", "agenda_theme", "relevance_band"]
LABEL_HEADERS = ["intent", "GT", "GA", "ADS", "opin", "merak", "gündem", "rel"]
SHEET_NAME = "ETIKETLEME"

_HDR_FILL = PatternFill("solid", fgColor="1F4E79")
_HDR_FONT = Font(color="FFFFFF", bold=True, size=11)
_KW_FILL = PatternFill("solid", fgColor="F2F2F2")
_EMPTY_FILL = PatternFill("solid", fgColor="FFC7CE")
_ALT_FILL = PatternFill("solid", fgColor="EAF1F8")
_THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def _dropdown_values(field: str, schema: Dict[str, list]) -> List[str]:
    return [str(v) for v in schema[field]] + ["?", "skip"]


# ── içerik: talimat / tanım / örnek metinleri ────────────────────────────

def _talimat_rows(kind: str, n_items: int) -> List[Tuple[str, str]]:
    ikinci = kind == "second"
    rows = [
        ("BU FORM NEDİR?",
         "Yapay zekâ modellerinin anahtar kelime kararlarını ölçmek için "
         "İNSAN REFERANSI oluşturuyoruz. Sizin etiketleriniz 'doğru cevap "
         "anahtarı' olacak; iki model bu cevaplara göre yarışacak. Model "
         "seçimini ve kalite kapılarını doğrudan sizin etiketleriniz "
         "belirleyecek."),
        ("NE YAPACAKSINIZ?",
         f"ETIKETLEME sekmesinde {n_items} kelime var. Her satırda 8 soru "
         "hücresi doldurulacak (hepsi açılır listeli — yazmak yok, "
         "seçmek var). Kırmızı görünen hücre = henüz boş. Kırmızı hücre "
         "kalmayınca iş bitmiştir."),
        ("ADIM ADIM",
         "1) Önce TANIMLAR sekmesini bir kez okuyun (5-10 dk). "
         "2) ÖRNEKLER sekmesine bakın — 6 kelime örnek olarak etiketlenmiş "
         "ve nedenleri yazılmış. "
         "3) ETIKETLEME sekmesine geçin; her kelime için kendinize şu "
         "soruyu sorun: 'Bunu Google'a yazan kişi NE İSTİYOR?' "
         "4) 8 hücreyi soldan sağa seçerek doldurun. "
         "5) Bitince dosyayı kaydedip geri gönderin."),
        ("EMİN DEĞİLSENİZ",
         "Kelimeyi Google'da aratıp çıkan sonuçlara bakmak SERBEST ve "
         "TAVSİYE EDİLİR (niyeti tahminle değil kanıtla etiketleyin). "
         "Hâlâ kararsızsanız hücrede '?' seçin — bu maddeler sonra "
         "birlikte karara bağlanır. '?' utanılacak şey değildir; yanlış "
         "kesin cevaptan iyidir."),
        ("'skip' NE ZAMAN?",
         "YALNIZ metin bozuksa (kırık karakterler, anlamsız dizi). "
         "DİKKAT: Konu dışı kelime 'skip' DEĞİLDİR! 'dolar kuru' gibi "
         "alakasız bir kelime tam da negatif örnektir — normal etiketlenir "
         "(tipik: ADS -1, GA 0, rel low). Modelin 'hayır' demeyi öğrenip "
         "öğrenmediğini bu kelimeler test eder."),
        ("KURALLAR",
         "• Her hücre dolu olmalı — boş hücre eksik iş sayılır ve sistem "
         "kabul etmez. "
         "• Kelime ve kimlik sütunları kilitlidir, değiştirmeyin. "
         "• Satır atlamak yok; sıra önemli değil, tamamlanması önemli. "
         "• Tahmini süre: 60-90 dakika. Tek oturmada bitirmek zorunda "
         "değilsiniz — kaydedin, sonra devam edin."),
    ]
    if ikinci:
        rows.insert(1, (
            "BAĞIMSIZLIK ŞARTI (ÖNEMLİ)",
            "Bu form İKİNCİ etiketçi formudur ve 30 kelime içerir. Birincil "
            "etiketçinin dosyasını AÇMAYIN, cevaplarını SORMAYIN, birlikte "
            "doldurmayın. İki bağımsız görüş, uyuşmazlıkları yakalamak için "
            "gerekli — kopya çekilirse ölçüm çöker. Uyuşmayan maddeler "
            "sonra ortak oturumda karara bağlanacak."))
    else:
        rows.insert(1, (
            "ÇİFT SÜTUNU NEDİR?",
            "'Çift' sütununda ✓ olan kelimeler ikinci bir kişi tarafından "
            "da BAĞIMSIZ etiketlenecek (kalite kontrolü). Sizin için fark "
            "yok: her kelimeyi normal etiketleyin. ✓'li maddelerde cevabı "
            "kimseyle tartışmayın."))
    return rows


_TANIMLAR: List[Tuple[str, str, str, str]] = [
    ("intent",
     "Bu kelimeyi arayan kişinin NİYETİ ne?",
     "transactional / commercial / informational / navigational / trend_worthy",
     "transactional: Bir şeyi EDİNMEK üzere ('satın al', 'indir', 'üyelik', "
     "'abonelik'). — commercial: Almadan önce ARAŞTIRIYOR ('en iyi X', "
     "'X karşılaştırma', 'hangisi'). — informational: Sadece ÖĞRENMEK "
     "istiyor ('nedir', 'nasıl', rehber). — navigational: Belirli bir "
     "siteye/uygulamaya GİTMEK istiyor ('hissefy giriş', 'kap resmi "
     "sitesi'). — trend_worthy: Güncel olaya/gündeme bağlı arama (haber, "
     "ani gelişme)."),
    ("GT",
     "Arayan kişi bir ÜRÜN/HİZMET edinmeye yakın mı? (0/1)",
     "1 = evet, yakın · 0 = hayır, salt bilgi/gündem",
     "intent'ten BAĞIMSIZ değerlendirin. 'portföy takip uygulaması öneri' "
     "bilgi sorusu gibi görünür ama araç arayışıdır → GT 1. 'bist "
     "şirketleri listesi' commercial görünebilir ama edinme niyeti yok → "
     "GT 0. Soru şu: bu kişiye bir araç/hizmet gösterilse ilgilenir mi?"),
    ("GA",
     "Kelime bizim ÜRÜN KATEGORİMİZLE ilgili bir arama mı? (0/1)",
     "1 = evet · 0 = hayır",
     "Ürün kategorisi: borsa/hisse ANALİZ ve TAKİP araçları, tarama, "
     "portföy takibi (hissefy'nin alanı). 'hisse tarama programı' → 1. "
     "'dolar kuru bugün', 'vadeli mevduat', 'borsa dolandırıcılığı "
     "haberi' → 0 (finans olabilir ama bizim kategori değil)."),
    ("ADS",
     "Bu kelimeye REKLAM verilse para eder mi?",
     "2 = hot_sale · 1 = lead · -1 = israf",
     "2: Tıklayan dönüşüme YAKIN — güçlü araç/edinme niyeti ('hisse "
     "tarama aracı', 'borsa uygulaması pro'). — 1: Reklam gösterilebilir, "
     "ilgi var ama dönüşüm dolaylı (çoğu araştırma araması). — -1: Reklam "
     "parası İSRAF — ödev, eğlence, şikayet, haber, alakasız aramalar. "
     "Kararsızsanız Google'da aratın: sonuçlar araç/uygulama gösteriyorsa "
     "arayan araç arıyor olabilir."),
    ("opin",
     "Bu konu sosyal medyada GÖRÜŞ/TARTIŞMA üretir mi? (0/1)",
     "1 = evet · 0 = hayır",
     "'X değer mi?', 'hangisi daha iyi?', yorum çeken temalar → 1. "
     "Kuru veri/tanım aramaları → 0."),
    ("merak",
     "MERAK veya KARŞILAŞTIRMA kancası var mı? (0/1)",
     "1 = evet · 0 = hayır",
     "'İnsanlar bunu merak eder, tıklar' dedirten açı: karşılaştırmalar, "
     "'gerçekten işe yarıyor mu' tarzı sorular → 1. Sıradan işlem/veri "
     "aramaları → 0."),
    ("gündem",
     "GÜNCEL gündeme/döneme mi bağlı? (0/1)",
     "1 = evet · 0 = hayır",
     "1: faiz kararı, rekor, halka arz sezonu gibi zamana bağlı konular. "
     "0: her zaman geçerli (evergreen) eğitim/araç konuları."),
    ("rel",
     "Kelime hissefy'nin ürün TEMASINA ne kadar yakın?",
     "high / mid / low",
     "high: doğrudan ürün alanı (hisse analiz/takip/tarama araçları). — "
     "mid: finans/borsa alanı ama ürün dışı (genel borsa haberi, endeks, "
     "şirket listesi). — low: alan dışı (dolar/altın/kredi/mevduat/"
     "kripto, eğlence, ödev)."),
    ("?",
     "Kararsız kaldığınızda",
     "her sütunda seçilebilir",
     "Google'a baktınız, yine emin değilsiniz → '?' seçin. Bu maddeler "
     "çözümleme turunda birlikte karara bağlanır."),
    ("skip",
     "YALNIZ bozuk/anlaşılamayan metin",
     "her sütunda seçilebilir",
     "Kırık kodlama, anlamsız karakter dizisi. Konu dışı kelime skip "
     "DEĞİLDİR — o normal etiketlenir (negatif sınıfın kendisidir)."),
]

# Örnek kelimeler 150'lik listede YOKTUR (etkilememeleri için).
_ORNEKLER: List[Tuple[str, ...]] = [
    ("portföy takip uygulaması indir", "transactional", "1", "1", "2",
     "0", "0", "0", "high",
     "Uygulama EDİNME eylemi ('indir') → transactional, GT 1. Ürün "
     "kategorimizin tam merkezi → GA 1, rel high. Dönüşüme çok yakın → "
     "ADS 2. Tartışma/merak/gündem kancası yok → üçü 0."),
    ("en iyi borsa analiz programı hangisi", "commercial", "1", "1", "2",
     "1", "1", "0", "high",
     "'En iyi... hangisi' = satın alma öncesi araştırma → commercial, "
     "GT 1. Kategori tam isabet → GA 1, high. Araç niyeti güçlü → ADS 2. "
     "'Hangisi' tartışma+karşılaştırma üretir → opin 1, merak 1."),
    ("bilanço nasıl okunur", "informational", "0", "0", "-1",
     "0", "0", "0", "mid",
     "Öğrenme sorusu → informational, GT 0. Borsa alanı ama analiz-aracı "
     "kategorisi değil → GA 0, rel mid. Reklam tıklaması israf → ADS -1."),
    ("faiz kararı sonrası hisseler", "trend_worthy", "0", "0", "-1",
     "1", "1", "1", "mid",
     "Güncel olaya bağlı → trend_worthy, gündem 1. Görüş ve merak üretir "
     "→ opin 1, merak 1. Ürün edinme niyeti yok → GT 0, ADS -1. Borsa "
     "alanı ama ürün dışı → GA 0, mid."),
    ("dolar mı altın mı yatırım", "informational", "0", "0", "-1",
     "1", "1", "0", "low",
     "Bizim kategori DEĞİL (döviz/altın) ama skip de değil → normal "
     "etiketlenir: GA 0, rel low, ADS -1. Karşılaştırma sorusu görüş ve "
     "merak çeker → opin 1, merak 1."),
    ("borsa terimleri sÃ¶zlÃ¼ÄŸÃ¼ (bozuk metin örneği)", "skip", "skip",
     "skip", "skip", "skip", "skip", "skip", "skip",
     "Metin kodlaması bozuk, ne yazdığı güvenilir okunamıyor → TÜM "
     "hücrelere skip. skip'in TEK meşru kullanımı budur."),
]


# ── form üretimi ─────────────────────────────────────────────────────────

def _write_text_sheet(ws, title: str, rows, widths, wrap_cols) -> None:
    ws.sheet_view.showGridLines = False
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.append([title])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14, color="1F4E79")
    ws.append([])
    for row in rows:
        ws.append(list(row))
        r = ws.max_row
        ws.cell(row=r, column=1).font = Font(bold=True)
        for c in range(1, len(row) + 1):
            cell = ws.cell(row=r, column=c)
            cell.alignment = Alignment(
                wrap_text=(c in wrap_cols), vertical="top")
        ws.row_dimensions[r].height = None  # otomatik


def _current_labels(md_path: Path, schema: Dict[str, list],
                    second: bool) -> Dict[str, Dict[str, object]]:
    text = md_path.read_text(encoding="utf-8")
    parsed = (parse_second_md if second else parse_primary_md)(text, schema)
    return {k: v["labels"] for k, v in parsed.items()}


def _cift_map(md_path: Path) -> Dict[str, str]:
    out = {}
    for line in md_path.read_text(encoding="utf-8").splitlines():
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        if cells and re.fullmatch(r"K\d{3}", cells[0]) and len(cells) >= 2:
            out[cells[0]] = cells[1]
    return out


def _label_to_cell(value) -> Optional[object]:
    if value in (None, "FILL", ""):
        return None
    return value  # int'ler int kalır, '?' / 'skip' / metin aynen


def build_workbook(json_path: Path, md_path: Path, out_path: Path,
                   *, second: bool) -> Dict[str, int]:
    """Formu güncel md durumundan üretir (kısmi etiketler taşınır)."""
    doc = json.loads(json_path.read_text(encoding="utf-8"))
    schema = doc["label_schema"]
    entries = doc["entries"]
    labels_now = _current_labels(md_path, schema, second)
    cift = {} if second else _cift_map(md_path)

    wb = Workbook()
    kind = "second" if second else "primary"
    _write_text_sheet(
        wb.active, "NASIL DOLDURULUR — LÜTFEN ÖNCE BUNU OKUYUN",
        _talimat_rows(kind, len(entries)), widths=(26, 120), wrap_cols={2})
    wb.active.title = "1_TALIMAT"
    ws_t = wb.create_sheet("2_TANIMLAR")
    _write_text_sheet(
        ws_t, "SÜTUN TANIMLARI — her sütun tek bir soruya cevaptır",
        [("Sütun: " + a, s, "Seçenekler: " + v, t)
         for a, s, v, t in _TANIMLAR],
        widths=(16, 42, 34, 110), wrap_cols={2, 3, 4})
    ws_o = wb.create_sheet("3_ORNEKLER")
    _write_text_sheet(
        ws_o, "ÇÖZÜLMÜŞ ÖRNEKLER (bu kelimeler sizin listenizde YOK)",
        [("Kelime", "intent", "GT", "GA", "ADS", "opin", "merak",
          "gündem", "rel", "Neden böyle etiketlendi?")] + list(_ORNEKLER),
        widths=(40, 14, 5, 5, 6, 6, 7, 8, 7, 100), wrap_cols={1, 10})
    for c in range(1, 11):
        cell = ws_o.cell(row=3, column=c)
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL

    ws = wb.create_sheet(SHEET_NAME)
    headers = (["item", "Çift", "Kelime"] if not second
               else ["item", "Kelime"]) + LABEL_HEADERS
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = _HDR_FONT
        cell.fill = _HDR_FILL
        cell.alignment = Alignment(horizontal="center")
    id_cols = 3 if not second else 2
    kw_col = id_cols  # kelime her iki düzende son kimlik sütunu
    widths = ([8, 6, 46] if not second else [8, 46]) + [15, 6, 6, 7, 7, 8, 9, 8]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = w

    for entry in entries:
        item = entry["item_id"]
        row: List[object] = [item]
        if not second:
            row.append(cift.get(item, ""))
        row.append(entry["keyword"])
        now = labels_now.get(item, {})
        for field in LABEL_FIELDS:
            row.append(_label_to_cell(now.get(field)))
        ws.append(row)
        r = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = _THIN
            if c <= id_cols:
                cell.fill = _KW_FILL
                cell.protection = Protection(locked=True)
            else:
                cell.protection = Protection(locked=False)
                cell.alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=kw_col).alignment = Alignment(wrap_text=True)

    n = len(entries)
    for off, field in enumerate(LABEL_FIELDS):
        col = get_column_letter(id_cols + 1 + off)
        dv = DataValidation(
            type="list",
            formula1='"' + ",".join(_dropdown_values(field, schema)) + '"',
            allow_blank=True, showErrorMessage=True,
            errorTitle="Geçersiz değer",
            error="Lütfen açılır listeden seçin (yazmayın).")
        ws.add_data_validation(dv)
        rng = f"{col}2:{col}{n + 1}"
        dv.add(rng)
        ws.conditional_formatting.add(
            rng, FormulaRule(formula=[f'ISBLANK({col}2)'],
                             fill=_EMPTY_FILL, stopIfTrue=False))
    ws.freeze_panes = ws.cell(row=2, column=id_cols + 1)
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.sort = False
    ws.protection.autoFilter = False

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"items": n, "labeled": sum(
        1 for v in labels_now.values()
        if all(x not in (None, "FILL") for x in v.values()))}


# ── geri aktarım: xlsx → md ──────────────────────────────────────────────

def _cell_to_token(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def import_workbook(xlsx_path: Path, json_path: Path, md_path: Path,
                    *, second: bool) -> Dict[str, object]:
    """Doldurulmuş formu doğrular ve md tablosunu yeniden yazar.

    Fail-closed: bilinmeyen değer, oynanmış kelime/kimlik, eksik/fazla
    satır → HİÇBİR ŞEY yazılmaz, tüm sorunlar listelenir.
    """
    doc = json.loads(json_path.read_text(encoding="utf-8"))
    schema = doc["label_schema"]
    entries = {e["item_id"]: e["keyword"] for e in doc["entries"]}
    cift = {} if second else _cift_map(md_path)

    wb = load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise AnnotationError(f"'{SHEET_NAME}' sekmesi yok: {xlsx_path}")
    ws = wb[SHEET_NAME]
    id_cols = 3 if not second else 2

    problems: List[str] = []
    rows: Dict[str, List[str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or row[0] is None:
            continue
        item = str(row[0]).strip()
        if not re.fullmatch(r"K\d{3}", item):
            continue
        if item in rows:
            problems.append(f"{item}: Excel'de yinelenen satır")
            continue
        keyword = _cell_to_token(row[id_cols - 1])
        if item not in entries:
            problems.append(f"{item}: bilinmeyen kimlik")
            continue
        if keyword != entries[item]:
            problems.append(
                f"{item}: kelime değiştirilmiş "
                f"({keyword!r} != {entries[item]!r})")
        rows[item] = [_cell_to_token(v)
                      for v in row[id_cols:id_cols + len(LABEL_FIELDS)]]
    missing = sorted(set(entries) - set(rows))
    if missing:
        problems.append(f"Excel'de eksik satırlar: {missing[:10]}"
                        + (" ..." if len(missing) > 10 else ""))
    if problems:
        raise AnnotationError(
            "EXCEL AKTARIMI REDDEDİLDİ:\n  - " + "\n  - ".join(problems))

    # md tablosunu yeniden kur — tablo ÜSTÜNDEKİ başlık metni korunur
    original = md_path.read_text(encoding="utf-8")
    lines = original.splitlines()
    table_start = next(i for i, ln in enumerate(lines)
                       if ln.strip().startswith("| item"))
    header_text = "\n".join(lines[:table_start])
    if second:
        head = "| item | Kelime | intent | GT | GA | ADS | opin | merak | gündem | rel |"
        sep = "|------|--------|--------|----|----|-----|------|-------|--------|-----|"
    else:
        head = "| item | Çift | Kelime | intent | GT | GA | ADS | opin | merak | gündem | rel |"
        sep = "|------|------|--------|--------|----|----|-----|------|-------|--------|-----|"
    out_lines = [header_text, head, sep]
    for item_id in sorted(entries):
        cells = [item_id]
        if not second:
            cells.append(cift.get(item_id, ""))
        cells.append(entries[item_id])
        cells.extend(rows[item_id])
        out_lines.append("| " + " | ".join(cells) + " |")
    new_md = "\n".join(out_lines) + "\n"

    # yazmadan ÖNCE üretilen tabloyu gerçek parser'dan geçir (fail-closed)
    parsed = (parse_second_md if second else parse_primary_md)(new_md, schema)
    # Satir DOLDURULMUS = FILL yok VE "?" yok ("?" cozumleme turunu bekler).
    # NOT: bu sayac "etiket rejimi" ile ILGISIZDIR; label_regime.FULLY_LABELED
    # ile karistirilmamasi icin anahtar adi `rows_complete`dir.
    filled = sum(1 for v in parsed.values()
                 if all(x not in ("FILL", "?")
                        for x in v["labels"].values()))
    unsure = sum(1 for v in parsed.values()
                 if any(x == "?" for x in v["labels"].values()))
    md_path.write_text(new_md, encoding="utf-8", newline="\n")
    return {"items": len(parsed), "rows_complete": filled,
            "fully_labeled": filled,  # geriye donuk uyumluluk (DEPRECATED)
            "with_question": unsure,
            "remaining": len(parsed) - filled}
