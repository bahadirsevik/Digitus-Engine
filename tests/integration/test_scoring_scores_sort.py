"""Integration tests for sorted scoring results."""

import io

from openpyxl import load_workbook

from app.database.models import EngineSelection, EngineStageResult


def _seed_score(
    make_keyword,
    make_keyword_score,
    *,
    workspace_id: int,
    run_id: int,
    keyword: str,
    **score_kwargs,
):
    kw = make_keyword(keyword, brand_profile_id=workspace_id)
    make_keyword_score(scoring_run_id=run_id, keyword_id=kw.id, **score_kwargs)
    return kw


def _scores(client, run_id, workspace_id, **params):
    response = client.get(
        f"/api/v1/scoring/runs/{run_id}/scores",
        params={"brand_profile_id": workspace_id, **params},
    )
    assert response.status_code == 200
    return response.json()


def test_xlsx_export_with_non_latin1_run_name_returns_200(
    client, make_workspace, make_scoring_run, make_keyword, make_keyword_score
):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(
        brand_profile_id=ws.id,
        name='Şişli İçerik "test"',
        status="scored",
        algorithm_version="v2",
    )
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="a", ads_score=1)

    response = client.get(
        f"/api/v1/scoring/runs/{run.id}/export/xlsx",
        params={"brand_profile_id": ws.id},
    )

    assert response.status_code == 200
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=\"")
    assert "filename*=UTF-8''" in disposition
    assert "%C5%9E" in disposition  # Ş
    assert disposition.count('"') == 2  # kullanici tirnaklari header'i bozmaz


def test_scores_sort_ads_score_desc_nulls_last(
    client, make_workspace, make_scoring_run, make_keyword, make_keyword_score
):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(brand_profile_id=ws.id, status="scored", algorithm_version="v2")
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="low", ads_score=1)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="none", ads_score=None)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="high", ads_score=3)

    data = _scores(client, run.id, ws.id, sort_by="ads_score", sort_dir="desc")

    assert data["total_scored"] == 3
    assert [s["keyword"] for s in data["scores"]] == ["high", "low", "none"]


def test_scores_sort_seo_and_social_score_desc(
    client, make_workspace, make_scoring_run, make_keyword, make_keyword_score
):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(brand_profile_id=ws.id, status="scored", algorithm_version="v2")
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="a", seo_score=4, social_score=2)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="b", seo_score=8, social_score=9)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="c", seo_score=1, social_score=5)

    seo = _scores(client, run.id, ws.id, sort_by="seo_score", sort_dir="desc")
    social = _scores(client, run.id, ws.id, sort_by="social_score", sort_dir="desc")

    assert [s["keyword"] for s in seo["scores"]] == ["b", "a", "c"]
    assert [s["keyword"] for s in social["scores"]] == ["b", "c", "a"]


def test_scores_sort_rank_asc_best_first_and_nulls_last(
    client, make_workspace, make_scoring_run, make_keyword, make_keyword_score
):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(brand_profile_id=ws.id, status="scored", algorithm_version="v2")
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="third", ads_rank=3)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="none", ads_rank=None)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="first", ads_rank=1)

    data = _scores(client, run.id, ws.id, sort_by="ads_rank", sort_dir="asc")

    assert [s["keyword"] for s in data["scores"]] == ["first", "third", "none"]


def test_scores_sort_stable_tiebreaker_is_repeatable(
    client, make_workspace, make_scoring_run, make_keyword, make_keyword_score
):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(brand_profile_id=ws.id, status="scored")
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="first-insert", ads_score=5)
    _seed_score(make_keyword, make_keyword_score, workspace_id=ws.id, run_id=run.id, keyword="second-insert", ads_score=5)

    first = _scores(client, run.id, ws.id, sort_by="ads_score", sort_dir="desc")
    second = _scores(client, run.id, ws.id, sort_by="ads_score", sort_dir="desc")

    assert [s["keyword"] for s in first["scores"]] == ["first-insert", "second-insert"]
    assert [s["keyword"] for s in second["scores"]] == ["first-insert", "second-insert"]


def test_scores_pagination_keeps_sort_without_duplicates(
    client, make_workspace, make_scoring_run, make_keyword, make_keyword_score
):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(brand_profile_id=ws.id, status="scored")
    for idx, score in enumerate([6, 5, 4, 3, 2, 1], start=1):
        _seed_score(
            make_keyword,
            make_keyword_score,
            workspace_id=ws.id,
            run_id=run.id,
            keyword=f"kw-{idx}",
            ads_score=score,
        )

    page_1 = _scores(client, run.id, ws.id, sort_by="ads_score", sort_dir="desc", limit=3, offset=0)
    page_2 = _scores(client, run.id, ws.id, sort_by="ads_score", sort_dir="desc", limit=3, offset=3)
    keywords = [s["keyword"] for s in page_1["scores"] + page_2["scores"]]

    assert page_1["total_scored"] == 6
    assert page_2["total_scored"] == 6
    assert keywords == ["kw-1", "kw-2", "kw-3", "kw-4", "kw-5", "kw-6"]
    assert len(keywords) == len(set(keywords))


def test_scores_sort_validation_rejects_invalid_values(client, make_workspace, make_scoring_run):
    ws = make_workspace(name="A", status="confirmed")
    run = make_scoring_run(brand_profile_id=ws.id, status="scored")

    bad_sort = client.get(
        f"/api/v1/scoring/runs/{run.id}/scores",
        params={"brand_profile_id": ws.id, "sort_by": "keyword"},
    )
    bad_dir = client.get(
        f"/api/v1/scoring/runs/{run.id}/scores",
        params={"brand_profile_id": ws.id, "sort_dir": "sideways"},
    )
    bad_limit = client.get(
        f"/api/v1/scoring/runs/{run.id}/scores",
        params={"brand_profile_id": ws.id, "limit": 501},
    )

    assert bad_sort.status_code == 422
    assert bad_dir.status_code == 422
    assert bad_limit.status_code == 422


def test_v3_scores_read_engine_selections_with_family_and_delivery_state(
    client, db_session, make_workspace, make_scoring_run, make_keyword,
    make_keyword_score,
):
    ws = make_workspace(name="V3", status="confirmed")
    run = make_scoring_run(
        brand_profile_id=ws.id,
        status="channel_assigned",
        algorithm_version="v3",
    )
    kw = make_keyword("spiral çelik boru", brand_profile_id=ws.id)
    make_keyword_score(
        scoring_run_id=run.id,
        keyword_id=kw.id,
        metrics_snapshot={"engine": "engine_v3_universe_v1", "volume": 500},
    )
    db_session.add(EngineStageResult(
        scoring_run_id=run.id,
        stage="family_a1",
        scope_type="family",
        scope_key="dictionary",
        payload={"families": [{"family_id": "boru", "family_name": "Çelik Borular"}]},
        model="test-model",
        prompt_sha="a" * 64,
        firm_block_sha256="b" * 64,
    ))
    db_session.add_all([
        EngineSelection(
            scoring_run_id=run.id, keyword_id=kw.id, channel="ADS",
            algorithm_rank=2, final_rank=1, exclude_reason=None,
            family_id="boru", pool_class="primary",
            scores={"Selection": 81.25, "Core": 0.9},
        ),
        EngineSelection(
            scoring_run_id=run.id, keyword_id=kw.id, channel="SEO",
            algorithm_rank=4, final_rank=None, exclude_reason="CAPACITY_LIMIT",
            family_id="boru", pool_class="secondary",
            scores={"final": 72.5, "bp": 0.8},
        ),
        EngineSelection(
            scoring_run_id=run.id, keyword_id=kw.id, channel="SOCIAL",
            algorithm_rank=3, final_rank=2, exclude_reason=None,
            priority="TREND_CONTENT", pool_class="TREND_CONTENT",
            scores={"social_score": 66.75},
        ),
    ])
    db_session.commit()

    data = _scores(client, run.id, ws.id)
    assert data["algorithm_version"] == "v3"
    assert data["total_scored"] == 1
    row = data["scores"][0]
    assert row["family_id"] == "boru"
    assert row["family_name"] == "Çelik Borular"
    assert float(row["ads_score"]) == 81.25
    assert float(row["seo_score"]) == 72.5
    assert float(row["social_score"]) == 66.75
    assert (row["ads_rank"], row["ads_final_rank"]) == (2, 1)
    assert row["seo_exclude_reason"] == "CAPACITY_LIMIT"
    assert row["social_priority"] == "TREND_CONTENT"

    export = client.get(
        f"/api/v1/scoring/runs/{run.id}/export/xlsx",
        params={"brand_profile_id": ws.id},
    )
    assert export.status_code == 200
    sheet = load_workbook(io.BytesIO(export.content), read_only=True).active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    values = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
    by_header = dict(zip(headers, values))
    assert by_header["Aile Adı"] == "Çelik Borular"
    assert by_header["ADS Skor"] == 81.25
    assert by_header["SEO Durum"] == "ELENDİ: CAPACITY_LIMIT"
    assert by_header["SOCIAL Öncelik"] == "TREND_CONTENT"
