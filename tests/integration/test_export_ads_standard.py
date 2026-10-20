# -*- coding: utf-8 -*-
"""ADS export'unda 15 KW teslim standardı raporu (plan_ads_15kw_standardi §4).

Collector fill_keyword_ids + standard_report taşır; csv/excel/docx tablo/satırları
KW / başlık / KW'li açıklama sayılarını, "standart altı" eksik sayılarını ve KW
listelerini gösterir. Yalnız olgu ve liste: yüzde/oran/puan YOK.
"""
import csv
import io
import os
import zipfile

from docx import Document
from openpyxl import load_workbook

from app.database.models import AdDescription, AdGenerationSet, AdGroup, AdHeadline
from app.exporters.csv_exporter import CsvExporter
from app.exporters.data_collector import ExportDataCollector
from app.exporters.docx_exporter import DocxExporter
from app.exporters.excel_exporter import ExcelExporter
from app.exporters.pdf_exporter import PdfExporter
from app.schemas.export import ExportSectionEnum

ADS = [ExportSectionEnum.ADS]
STATUS = "Standart altı: 12 başlık eksik, 3 KW'li açıklama eksik, CTA eksik: 2"
POOL = "Havuz 15 KW'den az (3 KW)"


def _make_run(db_session, make_workspace, make_scoring_run):
    ws = make_workspace(name="Ads Standard WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    gen_set = AdGenerationSet(scoring_run_id=run.id, version_number=1, status="active")
    db_session.add(gen_set)
    db_session.flush()
    group = AdGroup(
        scoring_run_id=run.id, generation_set_id=gen_set.id, group_name="Laptop Grubu",
        target_keyword_ids=[1, 2, 3],
        target_keywords=["laptop uno", "laptop dos", "laptop tres"],
        fill_keyword_ids=[3],
    )
    db_session.add(group)
    db_session.flush()
    for i, text in enumerate(["Laptop Uno Fırsatı", "Hemen İnceleyin", "Ücretsiz Kargo"]):
        db_session.add(AdHeadline(ad_group_id=group.id, headline_text=text,
                                  headline_type="benefit", sort_order=i))
    db_session.add(AdDescription(ad_group_id=group.id, sort_order=0,
                                 description_text="Laptop dos için hızlı teslimat.",
                                 description_type="value_prop"))
    db_session.commit()
    return run


def test_collector_carries_fill_ids_and_standard_report(
    db_session, make_workspace, make_scoring_run
):
    run = _make_run(db_session, make_workspace, make_scoring_run)
    group = ExportDataCollector(db_session).collect(run.id, ADS).ads.ad_groups[0]

    assert group.fill_keyword_ids == [3]
    r = group.standard_report
    assert (r.keyword_count, r.headline_count, r.kw_description_count) == (3, 3, 1)
    assert (r.missing_headlines, r.missing_kw_descriptions) == (12, 3)
    assert (r.cta_count, r.missing_cta) == (1, 2)          # "Hemen İnceleyin" CTA
    assert r.below_standard and r.pool_below_target
    assert r.keywords_in_headlines == ["laptop uno"]
    assert r.keywords_description_only == ["laptop dos"]
    assert r.keywords_unused == ["laptop tres"]
    assert r.fill_keywords == ["laptop tres"]


def test_csv_excel_docx_pdf_show_report_without_percentages(
    db_session, make_workspace, make_scoring_run, tmp_path
):
    run = _make_run(db_session, make_workspace, make_scoring_run)

    # CSV: mevcut sütunlar yerinde, yenileri sonda
    csv_path = str(tmp_path / "ads.zip")
    CsvExporter(db_session).export(run.id, ADS, csv_path)
    with zipfile.ZipFile(csv_path) as zf:
        rows = list(csv.reader(io.StringIO(zf.read("reklam_gruplari.csv").decode("utf-8-sig"))))
    header, row = rows[0], rows[1]
    assert header[:5] == ["Grup Adı", "Hedef Kelimeler", "Başlık Sayısı", "Açıklama Sayısı", "Negatif Sayısı"]
    cells = dict(zip(header, row))
    assert cells["KW Sayısı"] == "3" and cells["KW'li Açıklama Sayısı"] == "1"
    assert cells["Standart Durumu"] == f"{STATUS}; {POOL}"
    assert cells["Başlıkta Kullanılan KW"] == "laptop uno"
    assert cells["Yalnız Açıklamada KW"] == "laptop dos"
    assert cells["Hiç Kullanılmayan KW"] == "laptop tres"
    assert cells["15'e Tamamlamak İçin Eklenen KW"] == "laptop tres"

    # Excel: aynı sütunlar 'Reklam Grupları' sayfasında
    xlsx_path = str(tmp_path / "ads.xlsx")
    ExcelExporter(db_session).export(run.id, ADS, xlsx_path)
    sheet = [list(r) for r in load_workbook(xlsx_path)["Reklam Grupları"].iter_rows(values_only=True)]
    assert sheet[0] == header
    assert sheet[1][5:] == [3, 1, 1, f"{STATUS}; {POOL}", "laptop uno", "laptop dos",
                            "laptop tres", "laptop tres"]

    # Docx: grup başına olgu satırları, 15 KW "teslim standardı" (Google zorunluluğu değil)
    docx_path = str(tmp_path / "ads.docx")
    DocxExporter(db_session).export(run.id, ADS, docx_path)
    doc_text = "\n".join(p.text for p in Document(docx_path).paragraphs)
    assert "15 KW teslim standardı: 3 KW, 3 başlık, 1 KW'li açıklama, 1 CTA (kural tabanlı sayım)" in doc_text
    assert STATUS in doc_text and POOL in doc_text
    assert "Hiç kullanılmayan KW: laptop tres" in doc_text
    assert "15'e tamamlamak için eklenen KW: laptop tres" in doc_text
    assert "zorunlu" not in doc_text.lower()

    # PDF: üretimi bozulmaz
    pdf_path = str(tmp_path / "ads.pdf")
    PdfExporter(db_session).export(run.id, ADS, pdf_path)
    assert os.path.getsize(pdf_path) > 0

    # Yüzde/oran/puan yok
    standard_cells = row[5:] + [c for c in sheet[1][5:] if isinstance(c, str)]
    assert not any("%" in str(c) for c in standard_cells)
    assert "%" not in "\n".join(l for l in doc_text.splitlines() if "KW" in l)
