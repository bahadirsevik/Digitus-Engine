"""Export dağıtımı testleri (plan v4 — Dışa Aktarım sayfası kalkıyor).

Kapsam:
- Skorlanan-kelimeler XLSX: metrics_snapshot esas, sıfırlar korunur,
  sektör WorkspaceKeyword'den, legacy fallback etiketi
- /keywords/pool/export.xlsx: aktif havuz, izolasyon, route çakışmaz, boş 400
- Kanal havuz XLSX: zengin kolonlar, kanal izolasyonu, 400/404/409
- SEO run-sızıntı regresyonu (collect_run_seo_contents)
- NO_CONTENT: pre-check 422 + worker yarışı (dosya oluşmaz)
- ADS aktif-set: taslak varken dosyada aktif setin grupları
- Alan matrisi eksiksizliği: eski limitlerin ÜZERİNDE kayıt üç formatta TAM
- Tekil bölüm Excel'inde Özet sayfası yok
- Geçmiş: format+sections+limit
- Güvenlik: Excel formula injection inert; PDF'te <yatırım> & A & B kırılmaz
- Dashboard sections karışmaması: _latest_full_export yalnız 'all' job'u seçer
"""
from __future__ import annotations

import io
import uuid

import pytest
from openpyxl import load_workbook

from app.database.models import (
    AdDescription,
    AdGenerationSet,
    AdGroup,
    AdHeadline,
    ChannelPool,
    ContentOutput,
    ExportJob,
    IntentAnalysis,
    Keyword,
    KeywordScore,
    NegativeKeyword,
    SEOGeoContent,
    SocialCategory,
    SocialContent,
    SocialIdea,
    WorkspaceKeyword,
)


# ── Yardımcılar ─────────────────────────────────────────────────────────────


def _wb_from_response(resp):
    assert resp.status_code == 200, resp.text
    return load_workbook(io.BytesIO(resp.content))


def _pdf_text(path: str) -> str:
    """ReportLab PDF'inin content stream'lerini açıp metni çıkarır.

    Tam bir PDF parser değil ama ASCII metin assert'lerine yeter
    (codex post-review #6). ReportLab stream'leri Flate + ASCII85
    sarmalıyla yazar; her iki katman da çözülür.
    """
    import base64
    import re
    import zlib

    def _decode(block: bytes) -> bytes:
        try:
            return zlib.decompress(block)
        except Exception:
            pass
        try:
            body = block.strip()
            if not body.startswith(b"<~"):
                body = b"<~" + body
            raw = base64.a85decode(body, adobe=True)
            try:
                return zlib.decompress(raw)
            except Exception:
                return raw
        except Exception:
            return block

    with open(path, "rb") as f:
        data = f.read()
    chunks = [
        _decode(m.group(1))
        for m in re.finditer(rb"stream\r?\n(.*?)endstream", data, re.S)
    ]
    return b"\n".join(chunks).decode("latin-1", errors="ignore")


def _rows(ws):
    return [[cell.value for cell in row] for row in ws.iter_rows()]


def _make_seo_content(db, run_id, keyword_id, title, *, is_stale=False, legacy=False):
    """SEO içeriği: modern kayıt ContentOutput üstünden run'a bağlanır."""
    output_id = None
    if not legacy:
        output = ContentOutput(
            scoring_run_id=run_id,
            keyword_id=keyword_id,
            channel="SEO",
            content_type="blog_post",
            content_data={"title": title},
            is_stale=is_stale,
        )
        db.add(output)
        db.flush()
        output_id = output.id
    content = SEOGeoContent(
        content_output_id=output_id,
        keyword_id=keyword_id,
        title=title,
        intro_paragraph=f"{title} giriş paragrafı.",
        body_content=f"{title} gövde metni. İkinci cümle.",
        meta_description=f"{title} meta",
        word_count=500,
    )
    db.add(content)
    db.commit()
    return content


def _make_social_chain(db, run_id, n_contents=1, *, stale=False, content_kwargs=None):
    category = SocialCategory(
        scoring_run_id=run_id, category_name="Eğitim", is_stale=stale
    )
    db.add(category)
    db.flush()
    contents = []
    for i in range(n_contents):
        idea = SocialIdea(
            category_id=category.id,
            idea_title=f"Fikir {i + 1}",
            target_platform="instagram",
            content_format="reels",
            trend_alignment=0.5,
            is_stale=stale,
        )
        db.add(idea)
        db.flush()
        kwargs = dict(
            idea_id=idea.id,
            caption=f"Caption {i + 1} " + "uzun metin " * 30,
            hooks=[
                {"text": f"Hook {j} fikir {i + 1}", "style": "question"}
                for j in range(1, 5)
            ],
            hashtags=[f"tag{j}" for j in range(1, 9)],
            scenario=f"Senaryo {i + 1}",
            visual_suggestion=f"Görsel önerisi {i + 1}",
            video_concept=f"Video konsepti {i + 1}",
            industry_posting_suggestion="Salı 09:00",
            platform_notes="Reels 30sn",
            cta_text="Takip et",
            is_stale=stale,
        )
        kwargs.update(content_kwargs or {})
        content = SocialContent(**kwargs)
        db.add(content)
        contents.append(content)
    db.commit()
    return contents


def _make_ads_set(db, run_id, *, status="active", version=1, group_name="Grup A",
                  n_headlines=4, n_descriptions=3, n_negatives=12, n_keywords=8):
    gen_set = AdGenerationSet(
        scoring_run_id=run_id, version_number=version, status=status
    )
    db.add(gen_set)
    db.flush()
    group = AdGroup(
        scoring_run_id=run_id,
        generation_set_id=gen_set.id,
        group_name=group_name,
        target_keywords=[f"hedef kelime {i}" for i in range(1, n_keywords + 1)],
    )
    db.add(group)
    db.flush()
    for i in range(1, n_headlines + 1):
        db.add(AdHeadline(ad_group_id=group.id, headline_text=f"Başlık {i}"[:30]))
    for i in range(1, n_descriptions + 1):
        db.add(AdDescription(
            ad_group_id=group.id, description_text=f"Açıklama metni {i}"[:90]
        ))
    for i in range(1, n_negatives + 1):
        db.add(NegativeKeyword(ad_group_id=group.id, keyword=f"negatif {i}"))
    db.commit()
    return gen_set, group


# ── 1) Skorlanan-kelimeler XLSX: snapshot esas ──────────────────────────────


def test_scored_xlsx_uses_metrics_snapshot(
    client, db_session, make_workspace, make_keyword, make_scoring_run
):
    ws = make_workspace(name="Snap WS")
    # Keyword'e KASITLI farklı değerler — snapshot kazanmalı
    kw = make_keyword(
        "snapshot kelime", brand_profile_id=ws.id,
        monthly_volume=99999, trend_3m=77.0, competition_score=0.99,
    )
    wk = (
        db_session.query(WorkspaceKeyword)
        .filter(WorkspaceKeyword.keyword_id == kw.id)
        .one()
    )
    # Sektör workspace snapshot'ından okunmalı (global Keyword.sector değil)
    db_session.query(WorkspaceKeyword).filter(WorkspaceKeyword.id == wk.id).update(
        {"sector": "Enerji"}, synchronize_session=False
    )
    kw.sector = "YANLIŞ GLOBAL SEKTÖR"
    run = make_scoring_run(brand_profile_id=ws.id, status="scored", algorithm_version="v2")
    db_session.add(KeywordScore(
        scoring_run_id=run.id,
        keyword_id=kw.id,
        ads_score=10, seo_score=20, social_score=30,
        ads_rank=1, seo_rank=1, social_rank=1,
        metrics_snapshot={
            "wk_id": wk.id,
            "monthly_volume": 1200,
            "trend_3m": 25.0,
            "trend_12m": 0.0,  # SIFIR korunmalı ('or' kaybı regresyonu)
            "competition_score": 0.4,
        },
    ))
    db_session.commit()

    wb = _wb_from_response(client.get(
        f"/api/v1/scoring/runs/{run.id}/export/xlsx",
        params={"brand_profile_id": ws.id},
    ))
    rows = _rows(wb.active)
    header, row = rows[0], rows[1]
    idx = {name: i for i, name in enumerate(header)}
    assert row[idx["Aylık Hacim"]] == 1200
    assert row[idx["Trend 3M (%)"]] == 25.0
    assert row[idx["Trend 12M (%)"]] == 0.0
    assert row[idx["Rekabet Skoru"]] == 0.4
    assert row[idx["Sektör"]] == "Enerji"
    assert row[idx["Veri Kaynağı"]] == "workspace_snapshot"


def test_scored_xlsx_legacy_fallback_labelled(
    client, db_session, make_workspace, make_keyword, make_scoring_run
):
    ws = make_workspace(name="Legacy WS")
    kw = make_keyword("legacy kelime", brand_profile_id=ws.id, monthly_volume=750)
    run = make_scoring_run(brand_profile_id=ws.id, status="scored")
    db_session.add(KeywordScore(
        scoring_run_id=run.id, keyword_id=kw.id,
        ads_score=1, seo_score=1, social_score=1,
        metrics_snapshot=None,
    ))
    db_session.commit()

    wb = _wb_from_response(client.get(
        f"/api/v1/scoring/runs/{run.id}/export/xlsx",
        params={"brand_profile_id": ws.id},
    ))
    rows = _rows(wb.active)
    idx = {name: i for i, name in enumerate(rows[0])}
    assert rows[1][idx["Aylık Hacim"]] == 750
    assert rows[1][idx["Veri Kaynağı"]] == "legacy_keyword_fallback"


# ── 2) Kelime havuzu XLSX ───────────────────────────────────────────────────


def test_keyword_pool_xlsx_active_only_and_isolated(
    client, db_session, make_workspace, make_keyword
):
    ws_a = make_workspace(name="Havuz A")
    ws_b = make_workspace(name="Havuz B")
    make_keyword("aktif kelime", brand_profile_id=ws_a.id, monthly_volume=500)
    passive = make_keyword("pasif kelime", brand_profile_id=ws_a.id)
    passive.is_active = False
    make_keyword("baska workspace", brand_profile_id=ws_b.id)
    db_session.commit()

    wb = _wb_from_response(client.get(
        "/api/v1/keywords/pool/export.xlsx", params={"brand_profile_id": ws_a.id}
    ))
    assert wb.active.title == "Aktif Kelime Havuzu"
    keywords = [row[0] for row in _rows(wb.active)[1:]]
    assert keywords == ["aktif kelime"]


def test_keyword_pool_xlsx_empty_400_and_route_no_conflict(
    client, make_workspace
):
    ws = make_workspace(name="Boş WS")
    resp = client.get(
        "/api/v1/keywords/pool/export.xlsx", params={"brand_profile_id": ws.id}
    )
    assert resp.status_code == 400

    # Statik yol dinamik /{keyword_id} route'unu bozmadı
    detail = client.get("/api/v1/keywords/123", params={"brand_profile_id": ws.id})
    assert detail.status_code == 404  # 'pool' parse hatası (422) DEĞİL


# ── 3) Kanal havuz XLSX ─────────────────────────────────────────────────────


@pytest.fixture
def channel_pool_setup(db_session, make_workspace, make_keyword, make_scoring_run):
    ws = make_workspace(name="Kanal WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    kw1 = make_keyword("ads birinci", brand_profile_id=ws.id, monthly_volume=900)
    kw2 = make_keyword("ads ikinci", brand_profile_id=ws.id, monthly_volume=400)
    kw3 = make_keyword("seo kelimesi", brand_profile_id=ws.id, monthly_volume=100)
    for kw, ads, seo in [(kw1, 40, 10), (kw2, 30, 5), (kw3, 5, 35)]:
        db_session.add(KeywordScore(
            scoring_run_id=run.id, keyword_id=kw.id,
            ads_score=ads, seo_score=seo, social_score=1,
        ))
    db_session.add(ChannelPool(
        scoring_run_id=run.id, keyword_id=kw1.id, channel="ADS",
        final_rank=1, relevance_score=0.8, adjusted_score=32.0,
        is_strategic=True, pool_label=None,
    ))
    db_session.add(ChannelPool(
        scoring_run_id=run.id, keyword_id=kw2.id, channel="ADS",
        final_rank=2, relevance_score=0.5, adjusted_score=15.0,
        is_strategic=False, pool_label="rising_opportunity",
    ))
    db_session.add(ChannelPool(
        scoring_run_id=run.id, keyword_id=kw3.id, channel="SEO",
        final_rank=1, adjusted_score=35.0,
    ))
    db_session.add(IntentAnalysis(
        scoring_run_id=run.id, keyword_id=kw1.id, channel="ADS",
        intent_type="transactional", is_passed=True,
    ))
    db_session.commit()
    return ws, run, (kw1, kw2, kw3)


def test_channel_pool_xlsx_rich_columns_and_isolation(client, channel_pool_setup):
    ws, run, _ = channel_pool_setup
    wb = _wb_from_response(client.get(
        f"/api/v1/channels/runs/{run.id}/pools/ADS/export.xlsx",
        params={"brand_profile_id": ws.id},
    ))
    rows = _rows(wb.active)
    header = rows[0]
    assert header == ['Sıra', 'Keyword', 'Aylık Hacim', 'Kanal Skoru',
                      'Relevance', 'Adjusted Skor', 'Niyet', 'Stratejik', 'Etiket']
    body = rows[1:]
    assert [r[1] for r in body] == ["ads birinci", "ads ikinci"]  # SEO kelimesi YOK
    assert body[0][0] == 1 and body[0][2] == 900 and body[0][3] == 40.0
    assert body[0][4] == 0.8 and body[0][5] == 32.0
    assert body[0][6] == "transactional" and body[0][7] == "Evet"
    # openpyxl boş string hücreyi None okur — stratejik olmayanda işaret yok
    assert body[1][7] in ("", None) and body[1][8] == "rising_opportunity"


def test_channel_pool_xlsx_errors(client, channel_pool_setup, make_workspace):
    ws, run, _ = channel_pool_setup
    assert client.get(
        f"/api/v1/channels/runs/{run.id}/pools/BILINMEYEN/export.xlsx",
        params={"brand_profile_id": ws.id},
    ).status_code == 400

    other = make_workspace(name="Yabancı WS")
    assert client.get(
        f"/api/v1/channels/runs/{run.id}/pools/ADS/export.xlsx",
        params={"brand_profile_id": other.id},
    ).status_code == 404

    # SOCIAL havuzu hiç kurulmadı → boş 400
    assert client.get(
        f"/api/v1/channels/runs/{run.id}/pools/SOCIAL/export.xlsx",
        params={"brand_profile_id": ws.id},
    ).status_code == 400


def test_channel_pool_xlsx_stale_409(client, db_session, channel_pool_setup):
    ws, run, _ = channel_pool_setup
    run.channel_pool_policy_version = (ws.policy_version or 1) + 5
    db_session.commit()
    resp = client.get(
        f"/api/v1/channels/runs/{run.id}/pools/ADS/export.xlsx",
        params={"brand_profile_id": ws.id},
    )
    assert resp.status_code == 409
    assert resp.json()["detail"]["code"] == "POLICY_STALE"


# ── 4) SEO run-sızıntı regresyonu ───────────────────────────────────────────


def test_seo_collector_run_isolation_and_legacy_exclusion(
    db_session, make_workspace, make_keyword, make_scoring_run
):
    from app.exporters.data_collector import collect_run_seo_contents

    ws = make_workspace(name="Sızıntı WS")
    kw = make_keyword("ortak kelime", brand_profile_id=ws.id)
    run_a = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    run_b = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    for run in (run_a, run_b):
        db_session.add(ChannelPool(
            scoring_run_id=run.id, keyword_id=kw.id, channel="SEO", final_rank=1
        ))
    db_session.commit()

    _make_seo_content(db_session, run_a.id, kw.id, "Run A içeriği")
    _make_seo_content(db_session, run_b.id, kw.id, "Run B içeriği")
    # Legacy: content_output_id NULL — run ilişkisi kanıtlanamaz, dışlanır
    _make_seo_content(db_session, run_a.id, kw.id, "Legacy içerik", legacy=True)

    titles_a = [c.title for c in collect_run_seo_contents(db_session, run_a.id)]
    titles_b = [c.title for c in collect_run_seo_contents(db_session, run_b.id)]
    assert titles_a == ["Run A içeriği"]
    assert titles_b == ["Run B içeriği"]


# ── 5) NO_CONTENT: pre-check + worker yarışı ───────────────────────────────


def test_create_export_no_content_422(client, make_workspace, make_scoring_run):
    ws = make_workspace(name="Boş İçerik WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    resp = client.post(
        "/api/v1/export/",
        params={"brand_profile_id": ws.id},
        json={"scoring_run_id": run.id, "format": "excel", "sections": ["ads"]},
    )
    assert resp.status_code == 422
    assert resp.json()["detail"]["code"] == "NO_CONTENT"


def test_worker_no_content_race_marks_failed_without_file(
    client, db_session, make_workspace, make_scoring_run, monkeypatch
):
    from app.tasks import export_tasks
    from app.tasks.export_tasks import run_export_task

    ws = make_workspace(name="Yarış WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    contents = _make_social_chain(db_session, run.id, n_contents=1)

    monkeypatch.setattr(run_export_task, "delay", lambda export_id: None)
    resp = client.post(
        "/api/v1/export/",
        params={"brand_profile_id": ws.id},
        json={"scoring_run_id": run.id, "format": "excel", "sections": ["social"]},
    )
    assert resp.status_code == 200  # pre-check geçti (içerik vardı)
    export_id = resp.json()["export_id"]

    # Yarış: worker başlamadan içerik stale edilir
    contents[0].is_stale = True
    db_session.commit()

    export_tasks.execute_export_job(export_id)
    db_session.expire_all()
    job = db_session.query(ExportJob).filter(ExportJob.id == export_id).one()
    assert job.status == "failed"
    assert "NO_CONTENT" in (job.error_message or "")
    assert not job.filepath  # dosya OLUŞMADI


def test_count_matches_collector_when_content_output_stale(
    db_session, make_workspace, make_scoring_run
):
    """Codex post-review #1: collector bağlı ContentOutput.is_stale'i de
    filtreler — sayaç aynı filtreyi uygulamazsa pozitif sayıp boş dosya
    üretilebilirdi. Sayaç artık birebir aynı; include_stale bayrağı da
    collector'la aynı anlamı taşır."""
    from app.exporters.data_collector import count_exportable_content

    ws = make_workspace(name="Stale Output WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")

    # ADS: aktif set + grubu STALE ContentOutput'a bağla
    gen_set, group = _make_ads_set(db_session, run.id)
    output = ContentOutput(
        scoring_run_id=run.id, keyword_id=None, channel="ADS",
        content_type="ad_group", content_data={}, is_stale=True,
    )
    # keyword_id NOT NULL — gerçek bir keyword gerekir
    kw = Keyword(keyword="stale ads kw", normalized_keyword="stale ads kw")
    db_session.add(kw)
    db_session.flush()
    output.keyword_id = kw.id
    db_session.add(output)
    db_session.flush()
    group.content_output_id = output.id
    db_session.commit()

    assert count_exportable_content(db_session, run.id, "ads") == 0
    assert count_exportable_content(
        db_session, run.id, "ads", include_stale_content=True
    ) == 1

    # SOCIAL: içerik zinciri taze ama bağlı ContentOutput stale
    contents = _make_social_chain(db_session, run.id, n_contents=1)
    social_output = ContentOutput(
        scoring_run_id=run.id, keyword_id=kw.id, channel="SOCIAL",
        content_type="social_post", content_data={}, is_stale=True,
    )
    db_session.add(social_output)
    db_session.flush()
    contents[0].content_output_id = social_output.id
    db_session.commit()

    assert count_exportable_content(db_session, run.id, "social") == 0
    assert count_exportable_content(
        db_session, run.id, "social", include_stale_content=True
    ) == 1


def test_worker_post_export_recheck_deletes_file(
    client, db_session, make_workspace, make_scoring_run, monkeypatch
):
    """Codex post-review #1 (yarış): içerik EXPORT SIRASINDA stale olursa
    dosya-sonrası kesin doğrulama job'ı failed yapar ve dosyayı siler —
    bayat veri 'completed' rapor olarak teslim edilmez."""
    import os as _os

    from app.tasks import export_tasks
    from app.tasks.export_tasks import run_export_task

    ws = make_workspace(name="Mid-Export WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    contents = _make_social_chain(db_session, run.id, n_contents=1)
    content_id = contents[0].id

    monkeypatch.setattr(run_export_task, "delay", lambda export_id: None)
    resp = client.post(
        "/api/v1/export/",
        params={"brand_profile_id": ws.id},
        json={"scoring_run_id": run.id, "format": "excel", "sections": ["social"]},
    )
    assert resp.status_code == 200
    export_id = resp.json()["export_id"]

    produced = {}

    class MidExportStaleExporter:
        """Dosyayı üretir AMA üretim sırasında içerik stale olur."""

        class _Collector:
            include_stale_content = False

        def __init__(self, db):
            self.db = db
            self.data_collector = self._Collector()

        def export(self, scoring_run_id, sections, filepath):
            with open(filepath, "wb") as f:
                f.write(b"yarim kalmis rapor")
            produced["path"] = filepath
            self.db.query(SocialContent).filter(
                SocialContent.id == content_id
            ).update({"is_stale": True}, synchronize_session=False)
            self.db.commit()
            return filepath

    monkeypatch.setattr(
        export_tasks, "_get_exporter",
        lambda format_enum, db: MidExportStaleExporter(db),
    )

    export_tasks.execute_export_job(export_id)
    db_session.expire_all()
    job = db_session.query(ExportJob).filter(ExportJob.id == export_id).one()
    assert job.status == "failed"
    assert "NO_CONTENT" in (job.error_message or "")
    assert not job.filepath
    assert not _os.path.exists(produced["path"])  # dosya SİLİNDİ


def test_worker_in_place_regeneration_fails_content_changed(
    client, db_session, make_workspace, make_scoring_run, monkeypatch
):
    """Codex post-review-3 #1: SOCIAL regenerate YENİ satır üretmez — aynı
    ID'de caption/hooks yerinde güncellenir, yalnız regeneration_count artar.
    ID kümesi değişmediği için eski kimlik bunu kaçırırdı. Kimlik artık
    (id, content_regen, idea_regen) taşır ve kolon-bazlı sorguyla DB'den
    taze okunur; güncelleme AYRI session'dan yapılsa da yakalanır."""
    import os as _os

    from app.database.connection import SessionLocal
    from app.tasks import export_tasks
    from app.tasks.export_tasks import run_export_task

    ws = make_workspace(name="In-Place Regen WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    contents = _make_social_chain(db_session, run.id, n_contents=1)
    content_id = contents[0].id

    monkeypatch.setattr(run_export_task, "delay", lambda export_id: None)
    resp = client.post(
        "/api/v1/export/",
        params={"brand_profile_id": ws.id},
        json={"scoring_run_id": run.id, "format": "excel", "sections": ["social"]},
    )
    assert resp.status_code == 200
    export_id = resp.json()["export_id"]

    produced = {}

    class RegenerateDuringExportExporter:
        """Dosyayı üretir; üretim SIRASINDA aynı satır AYRI session'dan
        regenerate edilir (generator'ın gerçek deseni: in-place update +
        regeneration_count artışı). Satır sayısı ve ID kümesi SABİT kalır."""

        class _Collector:
            include_stale_content = False

        def __init__(self, db):
            self.db = db
            self.data_collector = self._Collector()

        def export(self, scoring_run_id, sections, filepath):
            with open(filepath, "wb") as f:
                f.write(b"eski caption ile rapor")
            produced["path"] = filepath
            other_session = SessionLocal()
            try:
                row = other_session.query(SocialContent).filter(
                    SocialContent.id == content_id
                ).one()
                row.caption = "REGENERATE EDİLMİŞ YENİ CAPTION"
                row.regeneration_count = (row.regeneration_count or 0) + 1
                other_session.commit()
            finally:
                other_session.close()
            return filepath

    monkeypatch.setattr(
        export_tasks, "_get_exporter",
        lambda format_enum, db: RegenerateDuringExportExporter(db),
    )

    export_tasks.execute_export_job(export_id)
    db_session.expire_all()
    job = db_session.query(ExportJob).filter(ExportJob.id == export_id).one()
    assert job.status == "failed"
    assert "CONTENT_CHANGED" in (job.error_message or "")
    assert not job.filepath
    assert not _os.path.exists(produced["path"])  # eski caption'lı dosya SİLİNDİ


def test_ads_identity_changes_on_set_activation(
    db_session, make_workspace, make_scoring_run
):
    """Codex post-review-2 #1: ADS set aktivasyonunda grup SAYISI aynı kalır
    ama kimlik kümesi (set_id, group_id) değişir — worker karşılaştırması
    bunu yakalar."""
    from app.exporters.data_collector import collect_exportable_identity

    ws = make_workspace(name="Set Aktivasyon WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    set1, _ = _make_ads_set(db_session, run.id, status="active", version=1,
                            group_name="V1 Grubu")
    set2, _ = _make_ads_set(db_session, run.id, status="draft", version=2,
                            group_name="V2 Grubu")

    before = collect_exportable_identity(db_session, run.id, "ads")
    # Aktivasyon: eski aktif archive edilir (partial unique index sırası),
    # sonra taslak aktif olur
    set1.status = "archived"
    db_session.flush()
    set2.status = "active"
    db_session.commit()
    after = collect_exportable_identity(db_session, run.id, "ads")

    assert len(before) == len(after) == 1  # sayı sabit
    assert before != after                 # kimlik değişti — sayaç bunu göremezdi


def test_worker_same_count_content_swap_fails_content_changed(
    client, db_session, make_workspace, make_scoring_run, monkeypatch
):
    """Codex post-review-2 #1: toplam SAYI aynı kalsa da içerik değişirse
    (bir içerik stale + eşit sayıda yeni içerik) kimlik kümesi farkı
    yakalanır — eski içeriği taşıyan dosya 'completed' olamaz, silinir."""
    import os as _os

    from app.tasks import export_tasks
    from app.tasks.export_tasks import run_export_task

    ws = make_workspace(name="Swap WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    # A: taze (export edilecek), C: başlangıçta STALE (kümede yok)
    content_a, content_c = _make_social_chain(db_session, run.id, n_contents=2)
    content_c.is_stale = True
    db_session.commit()
    a_id, c_id = content_a.id, content_c.id

    monkeypatch.setattr(run_export_task, "delay", lambda export_id: None)
    resp = client.post(
        "/api/v1/export/",
        params={"brand_profile_id": ws.id},
        json={"scoring_run_id": run.id, "format": "excel", "sections": ["social"]},
    )
    assert resp.status_code == 200
    export_id = resp.json()["export_id"]

    produced = {}

    class SwapDuringExportExporter:
        """Dosyayı üretir; üretim SIRASINDA A stale olur, C tazelenir —
        toplam sayı 1→1 sabit kalır ama içerik tamamen değişmiştir."""

        class _Collector:
            include_stale_content = False

        def __init__(self, db):
            self.db = db
            self.data_collector = self._Collector()

        def export(self, scoring_run_id, sections, filepath):
            with open(filepath, "wb") as f:
                f.write(b"eski icerikli rapor")
            produced["path"] = filepath
            self.db.query(SocialContent).filter(SocialContent.id == a_id).update(
                {"is_stale": True}, synchronize_session=False
            )
            self.db.query(SocialContent).filter(SocialContent.id == c_id).update(
                {"is_stale": False}, synchronize_session=False
            )
            self.db.commit()
            return filepath

    monkeypatch.setattr(
        export_tasks, "_get_exporter",
        lambda format_enum, db: SwapDuringExportExporter(db),
    )

    export_tasks.execute_export_job(export_id)
    db_session.expire_all()
    job = db_session.query(ExportJob).filter(ExportJob.id == export_id).one()
    assert job.status == "failed"
    assert "CONTENT_CHANGED" in (job.error_message or "")
    assert not job.filepath
    assert not _os.path.exists(produced["path"])  # eski içerikli dosya SİLİNDİ


# ── 6) ADS aktif-set + tekil bölümde Özet yok + eksiksizlik ─────────────────


def test_ads_export_uses_active_set_and_no_summary_sheet(
    db_session, make_workspace, make_scoring_run, tmp_path
):
    from app.exporters.excel_exporter import ExcelExporter
    from app.schemas.export import ExportSectionEnum

    ws = make_workspace(name="ADS Set WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    _make_ads_set(db_session, run.id, status="active", version=1,
                  group_name="Aktif Grup", n_negatives=12)
    _make_ads_set(db_session, run.id, status="draft", version=2,
                  group_name="Taslak Grup")

    path = str(tmp_path / "ads.xlsx")
    ExcelExporter(db_session).export(run.id, [ExportSectionEnum.ADS], path)
    wb = load_workbook(path)

    # Tekil içerik bölümü dosyasında Özet sayfası YOK (plan v4 tur-4 #6)
    assert "Özet" not in wb.sheetnames
    groups = [r[0] for r in _rows(wb["Reklam Grupları"])[1:]]
    assert groups == ["Aktif Grup"]  # taslak set inmez

    # Eksiksizlik: 12 negatif (eski docx limiti 10'du) tamamı ayrı sheet'te
    negatives = [r[1] for r in _rows(wb["Negatifler"])[1:]]
    assert len(negatives) == 12

    # Hedef kelimeler kesintisiz (8 kelime, eski [:5] kesmesi yok)
    target_cell = _rows(wb["Reklam Metinleri"])[1][1]
    assert "hedef kelime 8" in target_cell


def test_social_excel_full_field_matrix(
    db_session, make_workspace, make_scoring_run, tmp_path
):
    from app.exporters.excel_exporter import ExcelExporter
    from app.schemas.export import ExportSectionEnum

    ws = make_workspace(name="Social Matris WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    _make_social_chain(db_session, run.id, n_contents=11)  # eski PDF limiti 5'ti

    path = str(tmp_path / "social.xlsx")
    ExcelExporter(db_session).export(run.id, [ExportSectionEnum.SOCIAL], path)
    wb = load_workbook(path)
    rows = _rows(wb["Sosyal İçerikler"])

    # Plan §8: brief bazında gruplama + ana kelime/süre/uyarı/format detayı
    # kolonları eklendi (brief'siz eski veri "Eski (brief'siz)" grubunda)
    assert rows[0] == ['Brief', 'Fikir', 'Platform', 'Format', 'Ana Kelime',
                       "Hook'lar", 'Caption', 'Senaryo', 'Görsel Önerisi',
                       'Video Konsepti', 'CTA', 'Hashtagler',
                       'Sektör Paylaşım Önerisi', 'Platform Notları',
                       'İstenen Süre', 'Gerçek Süre', 'Süre Durumu', 'Uyarılar',
                       'Format Detayı']
    body = rows[1:]
    assert len(body) == 11  # kesme yok
    first = body[0]
    assert first[0] == "Eski (brief'siz)"  # brief_id yok -> legacy grup
    assert "Hook 4" in first[5]          # 4 hook'un TAMAMI (eski [:3] yok)
    assert len(first[6]) > 100           # caption kesilmedi (eski [:100])
    assert first[8].startswith("Görsel önerisi")
    assert first[9].startswith("Video konsepti")
    assert first[12] == "Salı 09:00"
    assert first[13] == "Reels 30sn"
    assert "#tag8" in first[11]          # 8 hashtag'in TAMAMI (eski [:5] yok)


def test_docx_and_pdf_write_full_content_bodies(
    db_session, make_workspace, make_keyword, make_scoring_run, tmp_path
):
    """Word/PDF artık özet satırı değil içerik gövdesi yazar; kesme yok.

    PDF ayrıca ReportLab mini-HTML güvenliği: '<yatırım>' ve 'A & B' içeren
    metinler export'u KIRMAZ (safe_paragraph_text escape'i).
    """
    from docx import Document

    from app.exporters.docx_exporter import DocxExporter
    from app.exporters.pdf_exporter import PdfExporter
    from app.schemas.export import ExportSectionEnum

    ws = make_workspace(name="Gövde WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")

    # SEO: 21 içerik (eski docx limiti 20, pdf limiti 10'du) + riskli karakterler
    for i in range(1, 22):
        kw = make_keyword(f"seo kelime {i}", brand_profile_id=ws.id)
        db_session.add(ChannelPool(
            scoring_run_id=run.id, keyword_id=kw.id, channel="SEO", final_rank=i
        ))
        db_session.commit()
        title = (
            f"İçerik {i} <yatırım> A & B ÇĞİÖŞÜ çğıöşü"
            if i == 1
            else f"İçerik {i}"
        )
        _make_seo_content(db_session, run.id, kw.id, title)

    _make_ads_set(db_session, run.id, group_name="Tam Grup", n_negatives=12)
    _make_social_chain(db_session, run.id, n_contents=11)

    docx_path = str(tmp_path / "rapor.docx")
    DocxExporter(db_session).export(
        run.id,
        [ExportSectionEnum.SEO_CONTENT, ExportSectionEnum.ADS, ExportSectionEnum.SOCIAL],
        docx_path,
    )
    doc_text = "\n".join(p.text for p in Document(docx_path).paragraphs)
    assert "İçerik 21" in doc_text            # SEO [:20] kesmesi kalktı
    assert "negatif 12" in doc_text           # negatif [:10] kesmesi kalktı
    assert "Fikir 11" in doc_text             # social [:10] kesmesi kalktı
    assert "Hook 4" in doc_text               # hook [:3] kesmesi kalktı
    assert "Senaryo 1" in doc_text            # yeni alan: scenario
    assert "Video Konsepti:" in doc_text      # yeni alan matrisi
    assert "#tag8" in doc_text                # hashtag [:10] kesmesi kalktı

    pdf_path = str(tmp_path / "rapor.pdf")
    pdf_exporter = PdfExporter(db_session)
    pdf_exporter.export(
        run.id,
        [ExportSectionEnum.SEO_CONTENT, ExportSectionEnum.ADS, ExportSectionEnum.SOCIAL],
        pdf_path,
    )
    # Eksiksizlik PDF İÇERİĞİNDEN kanıtlanır (codex post-review #6) —
    # boyut kontrolü yeterli değil. ASCII metinler stream'lerden okunur.
    pdf_text = _pdf_text(pdf_path)
    assert "seo kelime 21" in pdf_text        # SEO [:10] kesmesi kalktı
    assert "negatif 12" in pdf_text           # ADS negatiflerinin tamamı
    assert "Hook 4 fikir 11" in pdf_text      # 11. içerik + 4. hook (kesme yok)
    assert "Video konsepti 11" in pdf_text    # yeni alan matrisi PDF'te

    # WinAnsi Helvetica yerine gömülü Unicode TrueType ailesi kullanılır.
    # ToUnicode haritası PDF okuyucularının Türkçe karakterleri soru işaretine
    # çevirmeden göstermesini ve kopyalamasını sağlar.
    from reportlab.pdfbase import pdfmetrics

    assert pdf_exporter.styles["TurkishNormal"].fontName == pdf_exporter.font_name
    assert pdf_exporter.styles["TurkishHeading"].fontName == pdf_exporter.font_name_bold
    font = pdfmetrics.getFont(pdf_exporter.font_name)
    assert all(ord(char) in font.face.charToGlyph for char in "ÇĞİÖŞÜçğıöşü")
    with open(pdf_path, "rb") as pdf_file:
        pdf_bytes = pdf_file.read()
    assert pdf_bytes.count(b"/FontFile2") >= 2  # normal + bold font gömülü
    assert pdf_bytes.count(b"/ToUnicode") >= 2


# ── 7) Geçmiş: format + sections + limit ────────────────────────────────────


def test_export_history_carries_type_and_respects_limit(
    client, db_session, make_workspace, make_scoring_run
):
    ws = make_workspace(name="Geçmiş WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")
    for i in range(7):
        db_session.add(ExportJob(
            id=str(uuid.uuid4()),
            brand_profile_id=ws.id,
            scoring_run_id=run.id,
            status="completed",
            progress=100,
            format="docx" if i % 2 else "excel",
            sections=["ads"] if i % 2 else ["all"],
        ))
    db_session.commit()

    resp = client.get(
        f"/api/v1/export/run/{run.id}",
        params={"brand_profile_id": ws.id, "limit": 5},
    )
    assert resp.status_code == 200
    body = resp.json()
    assert len(body["exports"]) == 5
    for job in body["exports"]:
        assert job["format"] in {"docx", "excel"}
        assert job["sections"] in (["ads"], ["all"])


# ── 8) Güvenlik: Excel formula injection ────────────────────────────────────


def test_excel_formula_injection_neutralized(
    client, db_session, make_workspace, make_keyword
):
    ws = make_workspace(name="Güvenlik WS")
    make_keyword('=HYPERLINK("http://evil";"tıkla")', brand_profile_id=ws.id,
                 monthly_volume=100)
    make_keyword('\t=HYPERLINK("http://evil2")', brand_profile_id=ws.id,
                 monthly_volume=200)
    make_keyword("+CMD çalıştır", brand_profile_id=ws.id, monthly_volume=300)

    wb = _wb_from_response(client.get(
        "/api/v1/keywords/pool/export.xlsx", params={"brand_profile_id": ws.id}
    ))
    rows = _rows(wb.active)[1:]
    for row in rows:
        # Metin hücresi inert: ' önekiyle başlar; formül olarak yorumlanmaz
        assert row[0].startswith("'"), row[0]
        # Sayısal hücreler sayı kaldı
        assert isinstance(row[3], (int, float))


def test_safe_text_helpers_unit():
    from app.exporters.safe_text import safe_excel_text, safe_paragraph_text

    assert safe_excel_text('=SUM(A1)').startswith("'")
    assert safe_excel_text('\t=HYPERLINK("x")').startswith("'")  # tur-6 #4
    assert safe_excel_text(42) == 42
    assert safe_excel_text(None) is None
    assert safe_paragraph_text('<yatırım> & A B') == '&lt;yatırım&gt; &amp; A B'
    assert safe_paragraph_text('satır1\nsatır2') == 'satır1<br/>satır2'


# ── 9) Dashboard sections karışmaması ───────────────────────────────────────


def test_latest_full_export_ignores_section_only_jobs(
    db_session, make_workspace, make_scoring_run
):
    """Eski 'all' job + DAHA YENİ ADS-only job → tam-rapor akışı 'all' job'u
    görür; ADS-only tek başına 'tam rapor hazır' SAYILMAZ (tur-6 #1)."""
    from datetime import datetime, timedelta

    from app.core.dashboard.summary import _latest_export, _latest_full_export

    ws = make_workspace(name="Sections WS")
    run = make_scoring_run(brand_profile_id=ws.id, status="channel_assigned")

    full_id = str(uuid.uuid4())
    now = datetime.utcnow()
    db_session.add(ExportJob(
        id=full_id, brand_profile_id=ws.id, scoring_run_id=run.id,
        status="completed", progress=100, format="excel", sections=["all"],
        created_at=now - timedelta(hours=2),
    ))
    # 25 DAHA YENİ kanal export'u: eski son-20 taraması tam raporu kaçırırdı —
    # DB-level JSONB filtresi limitten bağımsız bulmalı (codex post-review #5)
    ads_ids = [str(uuid.uuid4()) for _ in range(25)]
    for i, ads_id in enumerate(ads_ids):
        db_session.add(ExportJob(
            id=ads_id, brand_profile_id=ws.id, scoring_run_id=run.id,
            status="completed", progress=100, format="docx", sections=["ads"],
            created_at=now - timedelta(minutes=25 - i),
        ))
    db_session.commit()

    full = _latest_full_export(db_session, ws.id, run.id)
    assert full is not None and full.id == full_id  # 25 ADS-only atlandı

    latest = _latest_export(db_session, ws.id, run.id)
    assert latest.id == ads_ids[-1]  # kart en-son job'u türüyle etiketler

    # ADS-only tek başına: full export bulunamaz
    db_session.query(ExportJob).filter(ExportJob.id == full_id).delete()
    db_session.commit()
    assert _latest_full_export(db_session, ws.id, run.id) is None
