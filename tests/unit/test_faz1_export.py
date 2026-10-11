import json

import pytest
from openpyxl import load_workbook

from app.core.benchmark import faz1_export as fx
from app.core.benchmark.reference_characterization import CharacterizationError, seal


def _art(config, run, gemini_usd, claude_usd=None):
    prov = {"gemini": {"requests_used": 10, "cost_settled_usd": gemini_usd}}
    if claude_usd is not None:
        prov["claude"] = {"requests_used": 4, "cost_settled_usd": claude_usd}
    return seal({"workspace_id": 21, "config_name": config, "scoring_run_id": run, "budget_summary": prov,
                 "wall_clock_s": {"setup": 1, "orchestration": 2, "content_ads": 3, "content_seo": 4,
                                  "content_social": 5},
                 "content": {"ads": {"n_groups": 7}, "seo": {"content_ids": [1, 2, 3]}, "social": {"n_items": 9}},
                 "content_errors": [], "created_at": "x"})


def _write(tmp_path, docs):
    paths = []
    for i, d in enumerate(docs):
        p = tmp_path / f"a{i}.json"
        p.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
        paths.append(str(p))
    return paths


def test_load_orders_and_verifies(tmp_path):
    paths = _write(tmp_path, [_art("Aday", 2, 0.1, 0.5), _art("G/G/G", 1, 0.3)])
    groups = fx.load_artifacts(paths)
    assert [a["config_name"] for a in groups[21]] == ["G/G/G", "Aday"]
    bad = json.loads(open(paths[0], encoding="utf-8").read())
    bad["scoring_run_id"] = 99
    open(paths[0], "w", encoding="utf-8").write(json.dumps(bad))
    with pytest.raises(CharacterizationError):
        fx.load_artifacts(paths)


def test_export_summary_numbers(tmp_path):
    groups = fx.load_artifacts(_write(tmp_path, [_art("Aday", 2, 0.1, 0.5), _art("G/G/G", 1, 0.3)]))
    out = tmp_path / "o"
    out.mkdir()
    fx.export_summary(groups, out)
    rows = list(load_workbook(out / "ozet.xlsx")["Özet"].values)
    assert rows[1][:3] == (21, "G/G/G", 1) and rows[1][6] == 0.3
    aday = rows[2]
    assert aday[4] == 0.1 and aday[5] == 0.5 and aday[6] == 0.6 and aday[12] == 7 and aday[13] == 3
    md = (out / "ozet.md").read_text(encoding="utf-8")
    assert "konsol faturası ayrı" in md and "claude=4" in md


def test_render_pools_rows():
    pools = {"Aday": [{"keyword_id": 1, "final_rank": 1}, {"keyword_id": 3, "final_rank": 2}],
             "G/G/G": [{"keyword_id": 2, "final_rank": 2}, {"keyword_id": 1, "final_rank": 1}]}
    side, ortak = fx.render_pools_rows(pools, {1: "a", 2: "b", 3: "c"})
    assert side == [["rank", "G/G/G", "Aday"], [1, "a", "a"], [2, "b", "c"]]
    assert ortak[0] == ["keyword", "in G/G/G", "in Aday"]
    assert ["a", "✓", "✓"] in ortak and ["b", "✓", "—"] in ortak and ["c", "—", "✓"] in ortak
    assert ortak[-1] == ["ortak (hepsinde)", 1]


def _it(label, ks="21|k", ok=True, viol=()):
    return {"key_label": label, "key_str": ks, "text_blocks": [("Başlık", "metin " + label)],
            "hard_rule_ok": ok, "hard_rule_violations": list(viol)}


def test_render_content_md_and_rows(tmp_path):
    items = {"G/G/G": {"21|k": _it("kredi")}, "Aday": {"21|k": _it("kredi", ok=False, viol=["x>30"]),
                                                      "21|z": _it("zam", "21|z")}}
    md = fx.render_content_md(items)
    assert md.index("### kredi") < md.index("### zam")
    assert "**G/G/G**\n\nBaşlık: metin kredi" in md
    assert "Hard-rule ihlalleri: x>30" in md
    zam = md[md.index("### zam"):]
    assert zam.index("**G/G/G**") < zam.index(fx.MISSING) < zam.index("**Aday**")
    rows = fx.content_rows(items)
    assert rows[0][0] == "key_label" and len(rows) == 4 and rows[2][5] == "x>30"
    fx._write_xlsx(tmp_path / "t.xlsx", {"Blog": rows, "RSA grupları": rows})
    assert load_workbook(tmp_path / "t.xlsx").sheetnames == ["Blog", "RSA grupları"]


def test_cli_no_db(tmp_path):
    paths = _write(tmp_path, [_art("G/G/G", 1, 0.3)])
    assert fx.main(["--artifacts", paths[0], "--out", str(tmp_path / "e"), "--no-db"]) == 0
    assert (tmp_path / "e" / "ozet.xlsx").exists()
