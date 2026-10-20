"""Kelime bazlı reklam görünümü (Excel sheet + CSV) — DB'siz birim testi."""
import csv
import os

from openpyxl import Workbook

from app.exporters.ads_standard_format import STANDARD_HEADERS, standard_lines
from app.exporters.csv_exporter import CsvExporter
from app.exporters.excel_exporter import ExcelExporter
from app.generators.ads.validators import build_standard_report
from app.schemas.export import AdGroupData, AdsData, DescriptionData, HeadlineData

HEADER = ['Kelime', 'Reklam Grubu', 'Açıklama 1', 'Açıklama 2',
          'Açıklama 3', 'Açıklama 4', 'Başlıklar']


def _ads() -> AdsData:
    a = AdGroupData(
        id=1, group_name='Grup A',
        target_keywords=['a1', 'a2', 'a3', 'ortak', 'a5'],
        headlines=[HeadlineData(headline_text='A-h1'), HeadlineData(headline_text='A-h2')],
        descriptions=[DescriptionData(description_text=f'A-desc-{i}') for i in range(1, 5)],
    )
    b = AdGroupData(
        id=2, group_name='Grup B',
        target_keywords=['ortak', 'b2'],
        headlines=[HeadlineData(headline_text='B-h1')],
        descriptions=[DescriptionData(description_text=f'B-desc-{i}') for i in range(1, 3)],
    )
    return AdsData(ad_groups=[a, b], total=2)


def _check_rows(rows):
    assert len(rows) == 7
    by_kw = {}
    for r in rows:
        by_kw.setdefault(r[0], []).append(r)
    a5 = by_kw['a5'][0]
    assert a5[1] == 'Grup A' and a5[5] == 'A-desc-4'
    ortak = {r[1]: r for r in by_kw['ortak']}
    assert len(by_kw['ortak']) == 2
    assert ortak['Grup A'][2:6] == [f'A-desc-{i}' for i in range(1, 5)]
    assert ortak['Grup A'][6] == 'A-h1 | A-h2'
    b = ortak['Grup B']
    assert b[2:4] == ['B-desc-1', 'B-desc-2']
    assert not b[4] and not b[5]
    assert b[6] == 'B-h1'


def test_excel_keyword_sheet():
    wb = Workbook()
    ExcelExporter(None)._add_ads_sheets(wb, _ads())
    assert wb.sheetnames.index('Kelime Bazlı Reklamlar') < wb.sheetnames.index('Reklam Metinleri')
    ws = wb['Kelime Bazlı Reklamlar']
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    assert rows[0] == HEADER
    _check_rows(rows[1:])
    assert ws.freeze_panes == 'A2'
    assert ws.auto_filter.ref


def test_csv_keyword_file_and_no_truncation(tmp_path):
    files = CsvExporter(None)._write_ads_csv(str(tmp_path), _ads())
    assert any(os.path.basename(f) == 'kelime_bazli_reklamlar.csv' for f in files)

    def read(name):
        with open(tmp_path / name, encoding='utf-8-sig', newline='') as f:
            return list(csv.reader(f))

    rows = read('kelime_bazli_reklamlar.csv')
    assert rows[0] == HEADER
    _check_rows(rows[1:])

    groups = read('reklam_gruplari.csv')
    assert 'a5' in groups[1][1]


# ==================== 15 KW teslim standardı sütunları (DB'siz) ====================

def _ads_with_report() -> AdsData:
    kws = ['a1', 'a2', 'a3']
    headlines = ['A-h1 a1', 'A-h2']
    descriptions = ['A-desc a2']
    report = build_standard_report(kws, [1, 2, 3], [3], headlines, descriptions)
    group = AdGroupData(
        id=1, group_name='Grup A', target_keywords=kws, fill_keyword_ids=[3],
        standard_report=report,
        headlines=[HeadlineData(headline_text=h) for h in headlines],
        descriptions=[DescriptionData(description_text=d) for d in descriptions],
    )
    return AdsData(ad_groups=[group, _ads().ad_groups[1]], total=2)   # 2. grup: raporsuz


def test_csv_and_excel_carry_standard_report_columns(tmp_path):
    ads = _ads_with_report()
    CsvExporter(None)._write_ads_csv(str(tmp_path), ads)
    with open(tmp_path / 'reklam_gruplari.csv', encoding='utf-8-sig', newline='') as f:
        rows = list(csv.reader(f))
    wb = Workbook()
    ExcelExporter(None)._add_ads_sheets(wb, ads)
    sheet = [list(r) for r in wb['Reklam Grupları'].iter_rows(values_only=True)]

    assert rows[0] == sheet[0]
    assert rows[0][5:] == STANDARD_HEADERS
    cells = dict(zip(rows[0], rows[1]))
    assert cells['CTA Sayısı'] == '0'
    assert cells['KW Sayısı'] == '3' and cells["KW'li Açıklama Sayısı"] == '1'
    assert cells['Standart Durumu'] == (
        "Standart altı: 13 başlık eksik, 3 KW'li açıklama eksik, CTA eksik: 3; Havuz 15 KW'den az (3 KW)")
    assert cells['Başlıkta Kullanılan KW'] == 'a1' and cells['Yalnız Açıklamada KW'] == 'a2'
    assert cells['Hiç Kullanılmayan KW'] == 'a3'
    assert cells["15'e Tamamlamak İçin Eklenen KW"] == 'a3'
    assert sheet[1][5:] == [3, 1, 0, cells['Standart Durumu'], 'a1', 'a2', 'a3', 'a3']
    # raporsuz (eski) grup: sütunlar boş, satır bozulmaz
    assert rows[2][5:] == [''] * len(STANDARD_HEADERS)
    # yüzde/oran/puan yok
    assert not any('%' in c for c in rows[1] + rows[2])


def test_standard_lines_are_facts_only():
    lines = standard_lines(_ads_with_report().ad_groups[0])
    assert lines[0] == "15 KW teslim standardı: 3 KW, 2 başlık, 1 KW'li açıklama, 0 CTA (kural tabanlı sayım)"
    assert "Standart altı: 13 başlık eksik, 3 KW'li açıklama eksik, CTA eksik: 3" in lines
    assert "Başlıkta kullanılan KW: a1" in lines
    assert "15'e tamamlamak için eklenen KW: a3" in lines
    assert not any('%' in line or 'zorunlu' in line.lower() for line in lines)
    assert standard_lines(_ads().ad_groups[0]) == []       # rapor yoksa satır yok
